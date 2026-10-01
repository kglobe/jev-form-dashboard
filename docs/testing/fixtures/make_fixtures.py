# 測試劇本（docs/testing/e2e-playbook.md）的固定資料。在這個資料夾執行 `python make_fixtures.py`，
# 重新產生所有 CSV 和 stores.xlsx，並印出劇本裡的預期數字；數字改了要回去更新劇本的「預期數字」。
import csv
import datetime as dt
from collections import defaultdict
from pathlib import Path

from openpyxl import Workbook

HERE = Path(__file__).parent
REGIONS = {'北區': 20000, '中區': 16000, '南區': 10000}  # 每月營收基數
PRODUCTS = {'P01': (0.5, 0.25, 100, 5), 'P02': (0.3, 0.40, 40, 15), 'P03': (0.2, 0.50, 20, 25)}  # 營收占比、1 月的毛利率（之後每月加 1 個百分點）、單價、日
FACTOR = {1: 1.0, 2: 1.1, 3: 1.2, 4: 1.3, 5: 1.4, 6: 1.5, 7: 1.6}
SALES_HEAD = ['日期', '區域', '產品代碼', '營收', '毛利', '數量']


def sales(months):
    rows = []
    for m in months:
        for region, base in REGIONS.items():
            # 北區 6 月驟降到 3 成，給異常檢查一個明顯的目標
            f = 0.3 if (region, m) == ('北區', 6) else FACTOR[m]
            for code, (share, margin, price, day) in PRODUCTS.items():
                revenue = round(base * share * f)
                rows.append([f'2026-{m:02d}-{day:02d}', region, code, revenue, round(revenue * (margin + 0.01 * (m - 1))), revenue // price])
    return rows


def write(name, head, rows):
    with open(HERE / name, 'w', encoding='utf-8', newline='') as f:
        w = csv.writer(f, lineterminator='\n')
        w.writerow(head)
        w.writerows(rows)


h1, july = sales(range(1, 7)), sales([7])
write('sales-h1.csv', SALES_HEAD, h1)
write('sales-07.csv', SALES_HEAD, july)
write('sales-07-missing-column.csv', [h for h in SALES_HEAD if h != '毛利'], [[v for i, v in enumerate(r) if i != 4] for r in july])
write('sales-07-text-number.csv', SALES_HEAD, [r[:3] + ['N/A'] + r[4:] if n == 0 else r for n, r in enumerate(july)])

returns = [[f'2026-{m:02d}-10', '北區', 'P01', 400 + 100 * (m - 1)] for m in range(2, 7)]
returns += [[f'2026-{m:02d}-18', '南區', 'P02', 150 + 50 * (m - 1)] for m in range(2, 7)]
returns.append(['2026-04-22', '中區', 'P09', 150])  # 主檔沒有 P09：查表時對不到
returns.sort()
write('returns.csv', ['日期', '區域', '產品代碼', '退貨金額'], returns)

# P02 重複：查表時用排最前面的「檯燈」
write('products.csv', ['產品代碼', '產品名稱', '類別'],
      [['P01', '電鍋', '家電'], ['P02', '檯燈', '家電'], ['P03', '筆記本', '文具'], ['P02', '檯燈（舊款）', '家電']])

# 標題列 4 欄，第 3 行只有 3 欄
write('bad-rows.csv', ['日期', '區域', '產品代碼', '營收'], [['2026-01-05', '北區', 'P01', '5000'], ['2026-01-15', '北區', 'P02']])

codes = list(PRODUCTS)
stores = {
    '台中店': ('營收', [(dt.date(2026, m, 10), codes[(m - 1) % 3], 2800 + 200 * m) for m in range(1, 7)]),
    '高雄店': ('金額', [(dt.date(2026, m, 20), codes[(m - 1) % 3], 1900 + 100 * m) for m in range(1, 7)]),
}
wb = Workbook()
wb.remove(wb.active)
for sheet, (money, rows) in stores.items():
    ws = wb.create_sheet(sheet)
    ws.append(['日期', '產品代碼', money])
    for d, code, v in rows:
        ws.append([d, code, v])
        ws.cell(ws.max_row, 1).number_format = 'yyyy-mm-dd'
wb.save(HERE / 'stores.xlsx')


def table(title, keys, values):
    print(f'\n{title}\n')
    print('| ' + ' | '.join(keys) + ' |')
    print('|' + '---|' * len(keys))
    for row in values:
        print('| ' + ' | '.join(str(v) for v in row) + ' |')


def by(rows, key, col):
    out = defaultdict(int)
    for r in rows:
        out[key(r)] += r[col]
    return dict(sorted(out.items()))


month = lambda r: r[0][:7]
for label, rows in [('1～6 月', h1), ('1～7 月（接上七月之後）', h1 + july)]:
    rev, gp, qty = by(rows, month, 3), by(rows, month, 4), by(rows, month, 5)
    table(f'銷售 {label}：每月', ['月份', '營收', '毛利', '數量', '毛利率'],
          [[k, rev[k], gp[k], qty[k], f'{gp[k] / rev[k]:.1%}'] for k in rev]
          + [['合計', sum(rev.values()), sum(gp.values()), sum(qty.values()), f'{sum(gp.values()) / sum(rev.values()):.1%}']])
    table(f'銷售 {label}：各區每月營收', ['月份', *REGIONS], [[k, *(by(rows, lambda r: (month(r), r[1]), 3)[(k, g)] for g in REGIONS)] for k in rev]
          + [['合計', *(by(rows, lambda r: r[1], 3)[g] for g in REGIONS)]])
    table(f'銷售 {label}：各產品營收', ['產品代碼', '營收', '數量'], [[k, v, by(rows, lambda r: r[2], 5)[k]] for k, v in by(rows, lambda r: r[2], 3).items()])
    print(f'\n筆數 {len(rows)}')

ret = by(returns, month, 3)
sales_m = by(h1, month, 3)
table('退貨：每月退貨金額與退貨率（對照來源，退貨率＝退貨金額 ÷ 銷售營收）', ['月份', '退貨金額', '銷售營收', '退貨率'],
      [[k, ret.get(k, '（空）'), v, f'{ret[k] / v:.1%}' if k in ret else '（空）'] for k, v in sales_m.items()])
overlap = [r for r in h1 if month(r) >= '2026-02']
table('退貨：各區退貨率（只算 2026-02～2026-06）', ['區域', '退貨金額', '銷售營收（2～6 月）', '退貨率'],
      [[g, by(returns, lambda r: r[1], 3).get(g, 0), by(overlap, lambda r: r[1], 3)[g],
        f'{by(returns, lambda r: r[1], 3).get(g, 0) / by(overlap, lambda r: r[1], 3)[g]:.1%}'] for g in REGIONS])
print(f'\n退貨筆數 {len(returns)}，退貨金額合計 {sum(r[3] for r in returns)}')
for sheet, (money, rows) in stores.items():
    print(f'{sheet}：{len(rows)} 筆，{money}合計 {sum(v for *_, v in rows)}')
