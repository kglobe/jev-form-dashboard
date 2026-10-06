# JEV 儀表板簡報，python-pptx 版（設計規格見 brief.md，文案取自 content.py）
# 產出: JEV儀表板_python-pptx.pptx（與本腳本同資料夾）
# 執行: python build_pptx.py
# 看圖: pwsh -File C:\Users\I26436\.claude\skills\ppt-master\scripts\pptx_export.ps1 JEV儀表板_python-pptx.pptx -OutDir qa-python-pptx -Anim
#
# 版面用 python-pptx；玻璃卡、陰影、光暈、羽化、連接線、轉場、morph 和所有動畫用 ppt-master 的 pptx_fx。
# 配色：MP072（深藍＋橘），改色後要先刪 assets/*.png 讓 icon 重轉。
import subprocess
import sys
import time
import urllib.parse
from pathlib import Path

from lxml import etree
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION, XL_LEGEND_POSITION, XL_TICK_MARK
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.oxml import parse_xml
from pptx.oxml.ns import nsdecls, qn
from pptx.util import Inches, Pt

T0 = time.time()
HERE = Path(__file__).parent
sys.path.insert(0, r"C:\Users\I26436\.claude\skills\ppt-master\scripts")
sys.path.insert(0, str(HERE))
sys.dont_write_bytecode = True  # 不在 docs/slides 留 __pycache__
import content as C  # noqa: E402
from pptx_fx import animate, cjk_font, connect, corner, deck, fill_alpha, glow, gradient, shadow, soft_edge, spacing, transition  # noqa: E402
from pptx_fx import fx as _fx  # noqa: E402

BG1, BG2 = "1B2632", "2C3B4D"
INK, MUTE, LINE, ACC, ACC_INK, CODE_BG = "EEE9DF", "C9C1B1", "C9C1B1", "FFB162", "1B2632", "111921"
ORB = {"orange": ("FFB162", 0.35), "rust": ("A35139", 0.55), "cream": ("EEE9DF", 0.12)}
ORB_POS = [  # 每頁 (cx, cy, d)：orange, rust, cream
    [(8.3, 1.0, 4.4), (6.4, 5.2, 3.6), (0.9, 5.1, 1.8)], [(9.7, 4.9, 3.6), (0.3, 0.3, 3.2), (5.0, 5.7, 1.6)],
    [(0.2, 5.3, 3.4), (9.8, 0.2, 3.6), (5.0, 0.0, 1.4)], [(5.0, 6.0, 4.0), (-0.2, 2.8, 3.0), (10.0, 2.8, 1.8)],
    [(9.9, 0.6, 3.4), (4.6, 6.0, 3.0), (0.0, 0.4, 1.6)], [(0.0, 0.2, 3.2), (10.0, 5.4, 3.8), (6.0, 5.9, 1.4)],
    [(10.1, 3.0, 3.6), (1.5, 6.0, 3.2), (4.8, -0.3, 1.6)], [(3.5, 6.1, 3.8), (9.9, 0.0, 3.2), (0.1, 2.6, 1.8)],
    [(5.0, -0.6, 4.2), (0.0, 5.6, 3.4), (9.9, 5.5, 1.8)], [(-0.3, 3.5, 3.4), (10.2, 1.8, 3.4), (7.0, 6.0, 1.4)],
    [(9.6, 5.8, 3.8), (2.0, -0.4, 3.0), (0.2, 4.8, 1.6)], [(8.0, 1.2, 4.2), (6.0, 5.4, 3.4), (1.0, 5.2, 1.8)],
]
TRANSITION = {1: "fade", 4: "wipe", 7: "split", 9: "zoom", 11: "push"}  # 其他頁 morph
FONT, MONO = "Noto Sans TC", "Consolas"
DISPLAY, H1, H2, BODY, CAPTION, CODE = 54, 26, 16, 14, 10.5, 11
MID, CENTER, RIGHT = MSO_ANCHOR.MIDDLE, PP_ALIGN.CENTER, PP_ALIGN.RIGHT


def clr(hex_, alpha=1.0):
    a = f'<a:alpha val="{round(alpha * 100000)}"/>' if alpha < 1 else ""
    return f'<a:srgbClr val="{hex_}">{a}</a:srgbClr>'


# ─── 形狀與文字 ────────────────────────────────────────────
def paint(sp, fill=None, fa=1.0, line=None, la=1.0, lw=0.75):
    # 拔掉 p:style，主題的預設陰影、線色才不會疊上來
    st = sp._element.find(qn("p:style"))
    if st is not None:
        sp._element.remove(st)
    if fill:
        sp.fill.solid()
        sp.fill.fore_color.rgb = RGBColor.from_string(fill)
        if fa < 1:
            fill_alpha(sp, fa)
    else:
        sp.fill.background()
    if line:
        sp.line.color.rgb = RGBColor.from_string(line)
        sp.line.width = Pt(lw)
        if la < 1:  # python-pptx 的線沒有透明度 API
            sp.line._get_or_add_ln().find(qn("a:solidFill"))[0].append(
                parse_xml(f'<a:alpha {nsdecls("a")} val="{round(la * 100000)}"/>'))
    else:
        sp.line.fill.background()
    return sp


def box(s, x, y, w, h, geom=MSO_SHAPE.RECTANGLE, **kw):
    return paint(s.shapes.add_shape(geom, Inches(x), Inches(y), Inches(w), Inches(h)), **kw)


def card(s, x, y, w, h, kind="glass", geom=MSO_SHAPE.ROUNDED_RECTANGLE):
    if kind == "code":
        sp = box(s, x, y, w, h, geom, fill=CODE_BG, fa=0.55, line=INK, la=0.25)
    elif kind == "hot":
        # 光暈畫在形狀後面，填色太透會整片透出來；用近不透明的深藍底，光暈只留在外框
        sp = box(s, x, y, w, h, geom, fill=BG2, fa=0.85, line=ACC, lw=1.25)
        glow(sp, 8, ACC, 0.45)
    else:
        sp = box(s, x, y, w, h, geom, fill=INK, fa=0.12, line=INK, la=0.35)
    if kind != "code":
        shadow(sp, blur=18, dist=6, angle=90, color="000000", alpha=0.25)
    if geom == MSO_SHAPE.ROUNDED_RECTANGLE:
        corner(sp, 0.12)
    return sp


def run(p, t, size, color, bold=False, track=None, font=FONT, fx=None):
    r = p.add_run()
    r.text = t
    r.font.size, r.font.bold = Pt(size), bold
    r.font.color.rgb = RGBColor.from_string(color)
    cjk_font(r.font, font)
    r.font._rPr.set("lang", "zh-TW")  # 標成中文，PowerPoint 才套避頭尾規則（「，」不會落在行首）
    if font == MONO:
        r.font._rPr.find(qn("a:ea")).set("typeface", FONT)
    if track:
        spacing(r.font, track)
    if fx:  # 文字陰影（H1、Display），fx="glow" 再加 ACC 光暈；effectLst 要排在 solidFill 後面
        g = f'<a:glow rad="{Pt(8)}">{clr(ACC, 0.45)}</a:glow>' if fx == "glow" else ""
        r.font._rPr.find(qn("a:solidFill")).addnext(parse_xml(
            f'<a:effectLst {nsdecls("a")}>{g}<a:outerShdw blurRad="{Pt(4)}" dist="{Pt(1.5)}" dir="5400000" '
            f'algn="ctr" rotWithShape="0">{clr("000000", 0.35)}</a:outerShdw></a:effectLst>'))
    return r


def frame(tf, anchor=None, inset=0.0):
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = Inches(inset)
    if anchor:
        tf.vertical_anchor = anchor
    return tf


def write(tf, paras, align=None, line=None):
    """paras: [(文字, 字級, 顏色, 粗體, 其他 run 參數), ...]，一個 tuple 一段。line 是行距倍數（換算成固定 pt）。"""
    for i, (t, size, color, bold, *kw) in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p._p.get_or_add_pPr().set("hangingPunct", "1")  # 標點可以懸在行尾，不會被擠到下一行開頭
        if align:
            p.alignment = align
        if line:
            p.line_spacing = Pt(size * line)
        for j, part in enumerate(t.split("\n")):
            if j:
                p.add_line_break()
            run(p, part, size, color, bold, **(kw[0] if kw else {}))


def text(s, x, y, w, h, lines, size, color, bold=False, align=None, anchor=None, line=None, **kw):
    tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    lines = [lines] if isinstance(lines, str) else lines
    write(frame(tb.text_frame, anchor), [(t, size, color, bold, kw) for t in lines], align, line)
    return tb


def label(sp, paras, inset=0.06):
    write(frame(sp.text_frame, MID, inset), paras, CENTER)
    return sp


def badge(s, x, y, d, n):
    return label(box(s, x, y, d, d, MSO_SHAPE.OVAL, line=ACC, lw=1.25),
                 [(str(n), CAPTION if d < 0.45 else H2, ACC, True)], inset=0)


def dot(s, x, y, d, color, halo=False):
    sp = box(s, x, y, d, d, MSO_SHAPE.OVAL, fill=color)
    if halo:
        glow(sp, 6, ACC, 0.6)
    return sp


def wire(s, x1, y1, x2, y2, begin=None, end=None):
    """直線連接線，座標自己給；begin/end 是 (形狀, 連接點編號)，用形狀幾何真正的連接點編號。
    pptx_fx.connect 依外框算 4 個點，六角形的連接點不在那 4 個位置，所以六角形改用這個。"""
    c = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    c.line.color.rgb = RGBColor.from_string(LINE)
    c.line.width = Pt(1.25)
    c.line._get_or_add_ln().append(parse_xml(f'<a:tailEnd {nsdecls("a")} type="triangle"/>'))
    cnv = c._element.nvCxnSpPr.cNvCxnSpPr
    for tag, ref in (("a:stCxn", begin), ("a:endCxn", end)):
        if ref:
            cnv.append(parse_xml(f'<{tag} {nsdecls("a")} id="{ref[0].shape_id}" idx="{ref[1]}"/>'))
    return c


def edge(s, a, a_site, b, b_site):
    return connect(s, a, a_site, b, b_site, kind="straight", color=LINE, width=1.25)


def icons(names, hot):
    """SVG 轉透明底 PNG（Chrome headless 截 data URL）。一個要好幾秒，所以並行跑；
    ponytail: 轉過就沿用，快取只看檔名，改顏色要先刪 assets/*.png"""
    pngs = [HERE / "assets" / f"{n}.png" for n in names]
    procs = []
    for i, png in enumerate(pngs):
        if png.exists():
            continue
        src = png.with_suffix(".svg").read_text(encoding="utf-8").replace(
            'fill="currentColor"', f'width="256" height="256" fill="#{ACC if i == hot else INK}"')
        procs.append(subprocess.Popen([r"C:\Program Files\Google\Chrome\Application\chrome.exe", "--headless=new",
                                       "--disable-gpu", "--hide-scrollbars", "--default-background-color=00000000",
                                       "--window-size=256,256", f"--screenshot={png}",
                                       "data:image/svg+xml," + urllib.parse.quote(src)],
                                      stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL))
    for p in procs:
        p.wait(120)
    assert all(p.exists() for p in pngs), "icon 轉檔失敗"
    return pngs


def fx(shapes, effect, trigger="with", dur=400, **kw):  # brief：動畫沒寫時間就 400ms
    return _fx(shapes, effect, trigger, dur, **kw)


# ─── 投影片骨架 ─────────────────────────────────────────────
prs = deck()
prs.core_properties.author = "Wayne Chang"
prs.core_properties.title = "JEV 只做判斷，其他交給規則"
slides = []


def new_slide(page, title=None, footer=C.FOOTER_LABEL):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    slides.append(s)
    gradient(s.background, BG1, BG2, 315)  # python-pptx 的 315 = OOXML lin ang 45°，左上到右下
    for name, (cx, cy, d) in zip(ORB, ORB_POS[page - 1]):
        color, a = ORB[name]
        sp = box(s, cx - d / 2, cy - d / 2, d, d, MSO_SHAPE.OVAL, fill=color, fa=a)
        soft_edge(sp, d * 72 * 0.25)
        sp.name = f"!!orb-{name}"  # morph 靠名字配對
    kind = TRANSITION.get(page, "morph")
    transition(s, kind, dir={"wipe": "r", "push": "u"}.get(kind))  # wipe 從左、push 往上
    if title:
        text(s, 0.5, 0.4, 9.0, 0.9, title, H1, INK, True, anchor=MID, fx="shadow")
    if page > 1:
        if footer:
            text(s, 0.5, 5.0 if page < 12 else 4.95, 5.0 if page < 12 else 7.5, 0.25, footer, CAPTION, MUTE)
        text(s, 8.5, 5.0, 1.0, 0.25, f"{page:02d}", CAPTION, MUTE, align=RIGHT)
    return s


def notes(s, t):
    s.notes_slide.notes_text_frame.text = t


links = []  # (形狀, 目標頁)，等全部頁面建完再設


# ─── 01 封面（O1） ─────────────────────────────────────────
d = C.COVER
s = new_slide(1)
c1 = card(s, 0.35, 1.2, 7.2, 3.55)
kick = text(s, 0.7, 1.45, 6.6, 0.35, d["kicker"], H2, ACC, True, track=3)
ttl = text(s, 0.7, 1.85, 6.6, 1.75, d["title"], DISPLAY, INK, True, line=1.1, fx="shadow")
sub = text(s, 0.7, 3.7, 6.6, 0.4, d["subtitle"], H2, MUTE, True)
meta = text(s, 0.7, 4.2, 6.6, 0.3, d["meta"], CAPTION, MUTE)
animate(s, fx([c1, kick], "fade", "after", 500), fx(ttl, "zoom", "after", 600),
        fx(sub, "wipe", "after", side="l"), fx(meta, "fade", "after"))
notes(s, d["notes"])

# ─── 02 三個問題（G1） ─────────────────────────────────────
d = C.PROBLEMS
s = new_slide(2, d["title"])
seq = []
for i, (head, desc) in enumerate(d["cards"]):
    x, y, hot = 0.5 + i * 3.1, 1.6, i == 2
    cd = card(s, x, y, 2.8, 2.9, "hot" if hot else "glass")
    parts = [cd, text(s, x + 0.25, y + 0.25, 1.0, 0.4, f"{i + 1:02d}", H2, ACC, True),
             text(s, x + 0.25, y + 0.8, 2.3, 0.45, head, H2, INK, True),
             text(s, x + 0.25, y + 1.35, 2.3, 1.35, desc, BODY, MUTE, line=1.3)]
    seq.append(fx(parts, "fly", "click", side="b"))
    if hot:
        seq.append(fx(cd, "teeter", "after"))
animate(s, *seq)
notes(s, d["notes"])

# ─── 03 分工（C1，關鍵頁） ─────────────────────────────────
d = C.SPLIT
s = new_slide(3, d["title"])
seq = []
for (head, items), x, hc, dc, side in ((d["left"], 0.5, ACC, ACC, "l"), (d["right"], 5.15, INK, MUTE, "r")):
    parts = [card(s, x, 1.55, 4.35, 2.75), text(s, x + 0.3, 1.75, 3.8, 0.4, head, H2, hc, True)]
    for j, t in enumerate(items):
        y = 2.3 + j * 0.48
        parts += [dot(s, x + 0.3, y + 0.15, 0.1, dc), text(s, x + 0.5, y, 3.6, 0.4, t, BODY, INK, anchor=MID)]
    seq.append(fx(parts, "wipe", "click", side=side))
bottom = text(s, 0.5, 4.5, 9.0, 0.4, d["bottom"], BODY, ACC, True)
animate(s, *seq, fx(bottom, "zoom", "click"), fx(bottom, "pulse", "after"))
notes(s, d["notes"])

# ─── 04 全流程（P3，原生形狀加連接線） ─────────────────────
d = C.FLOW
s = new_slide(4, d["title"])
xs = [0.5, 2.867, 5.233, 7.6]
nodes = []
for i, (name, jev, target) in enumerate(d["nodes"]):
    x, y = (xs[i], 1.65) if i < 4 else (xs[7 - i], 3.25)
    if len(name) > 10:  # ponytail: 只有「組裝 mapping 與版面」超寬，在最後一個空白換行免得「面」落單；標籤變長再改成量字寬
        name = "\n".join(name.rsplit(" ", 1))
    n = label(card(s, x, y, 1.9, 0.8, "hot" if jev else "glass"), [(f"{i + 1:02d}", CAPTION, ACC, False), (name, 13, INK, True)])
    nodes.append(n)
    if target:
        links.append((n, target))
# 連接點：0 上、1 左、2 下、3 右。第 1 列右接左，4→5 下接上，第 2 列由右往左，左接右
sites = [(3, 1)] * 3 + [(2, 0)] + [(1, 3)] * 3
lines = [edge(s, a, sa, b, sb) for (a, b), (sa, sb) in zip(zip(nodes, nodes[1:]), sites)]
text(s, 0.5, 4.4, 9.0, 0.35, d["caption"], BODY, MUTE)
jev = [n for n, (_, j, _) in zip(nodes, d["nodes"]) if j]
animate(s, fx(nodes[0], "fade", "after"), *[fx([ln, n], "fade", "after", 300) for ln, n in zip(lines, nodes[1:])],
        fx(jev, "pulse", "click"))
notes(s, d["notes"])

# ─── 05 結構化題目（K3，程式碼代替圖表） ───────────────────
d = C.QUESTIONS
s = new_slide(5, d["title"])
ask = [card(s, 0.5, 1.55, 5.7, 1.85, "code"), text(s, 0.7, 1.65, 0.5, 0.25, d["ask_label"], CAPTION, ACC),
       text(s, 0.7, 1.9, 5.3, 1.45, d["ask_code"], CODE, INK, font=MONO)]
ans = [card(s, 0.5, 3.6, 5.7, 1.25, "code"), text(s, 0.7, 3.7, 0.5, 0.25, d["answer_label"], CAPTION, ACC),
       text(s, 0.7, 3.95, 5.3, 0.8, d["answer_code"], CODE, INK, font=MONO)]
spark = dot(s, 3.35 - 0.07, 3.4 - 0.07, 0.14, ACC, halo=True)  # 問卡底部中央
points = []
for i, (t, y) in enumerate(zip(d["points"], (1.6, 2.65, 3.7))):
    points.append(fx([badge(s, 6.5, y, 0.36, i + 1), text(s, 7.0, y - 0.05, 2.5, 0.95, t, BODY, INK)], "float", "click"))
animate(s, fx(ask, "fade", "after"), fx(spark, "appear", "click"), fx(spark, "path", "with", 800, by=(0, 0.2)),
        fx(ans, "split", "after"), fx(spark, "fade_out", "after"), *points)
notes(s, d["notes"])

# ─── 06 防呆（P2，決策流程） ───────────────────────────────
d = C.GUARD
s = new_slide(6, d["title"])
start = label(card(s, 0.5, 2.0, 1.6, 0.9), [(d["start"], BODY, INK, False)])
chk1 = label(card(s, 2.5, 1.95, 2.3, 1.0, geom=MSO_SHAPE.HEXAGON), [(d["check1"], 13, INK, False)], inset=0.1)
chk2 = label(card(s, 5.2, 1.95, 2.5, 1.0, geom=MSO_SHAPE.HEXAGON), [(d["check2"], 13, INK, False)], inset=0.1)
accept = label(card(s, 8.0, 2.0, 1.5, 0.9), [(d["accept"], BODY, INK, False)])
ignore = label(card(s, 2.9, 3.55, 4.4, 0.75, "hot"), [(d["ignore"], H2, INK, True)])
# 六角形連接點：0 右中、1 右下頂點、2 左下頂點、3 左中、4 左上、5 右上；圓角矩形 0 上、1 左、2 下、3 右
l_start = wire(s, 2.1, 2.45, 2.5, 2.45, (start, 3), (chk1, 3))
l_no1 = wire(s, 4.8, 2.45, 5.2, 2.45, (chk1, 0), (chk2, 3))
l_no2 = wire(s, 7.7, 2.45, 8.0, 2.45, (chk2, 0), (accept, 1))
l_yes1 = wire(s, 3.65, 2.95, 3.65, 3.55)  # 六角形底邊中點沒有連接點，「是」兩條不綁
l_yes2 = wire(s, 6.45, 2.95, 6.45, 3.55)
tag = lambda x, y, t: text(s, x, y, 0.3, 0.25, t, CAPTION, MUTE, align=CENTER)
t_no1, t_no2 = tag(4.85, 2.15, d["no"]), tag(7.7, 2.15, d["no"])
t_yes1, t_yes2 = tag(3.72, 3.12, d["yes"]), tag(6.52, 3.12, d["yes"])
bottom = text(s, 0.5, 4.55, 9.0, 0.35, d["bottom"], BODY, ACC, True)
animate(s, fx(start, "fade", "click"),
        fx(l_start, "wipe", "click", side="l"), fx(chk1, "fade"),
        fx(l_yes1, "wipe", "click", side="t"), fx(t_yes1, "appear"), fx(ignore, "dissolve"),
        fx(l_no1, "wipe", "click", side="l"), fx(t_no1, "appear"), fx(chk2, "fade"),
        fx(l_yes2, "wipe", "click", side="t"), fx(t_yes2, "appear"),
        fx(l_no2, "wipe", "click", side="l"), fx(t_no2, "appear"), fx(accept, "fade"),
        fx(bottom, "fade", "click"))
notes(s, d["notes"])

# ─── 07 合法組合表（C2，表格物件） ─────────────────────────
d = C.LEGAL
s = new_slide(7, d["title"])
rows = [d["header"]] + d["rows"]
gf = s.shapes.add_table(len(rows), 2, Inches(0.5), Inches(1.55), Inches(9.0), Inches(3.1))
tbl = gf.table
tbl.first_row, tbl.horz_banding = True, False
gf._element.graphic.graphicData.tbl.tblPr.find(qn("a:tableStyleId")).text = "{2D5ABB26-0587-4C30-8999-92F81FD0307C}"  # No Style, No Grid
tbl.columns[0].width, tbl.columns[1].width = Inches(4.0), Inches(5.0)
hline = f'<a:solidFill>{clr(INK, 0.25)}</a:solidFill>'
for r, cells in enumerate(rows):
    tbl.rows[r].height = Inches(3.1 / len(rows))
    for c, t in enumerate(cells):
        cell = tbl.cell(r, c)
        cell.margin_left, cell.margin_right = Inches(0.2), Inches(0.1)
        cell.margin_top = cell.margin_bottom = Inches(0.03)
        cell.vertical_anchor = MID
        write(cell.text_frame, [(t, H2, ACC_INK, True) if r == 0 else (t, BODY, INK, False)])
        cell.fill.solid()
        cell.fill.fore_color.rgb = RGBColor.from_string(ACC if r == 0 else INK)
        tcPr = cell._tc.get_or_add_tcPr()
        if r:  # 內文列白 8% / 14% 交替
            tcPr.find(qn("a:solidFill"))[0].append(parse_xml(f'<a:alpha {nsdecls("a")} val="{8000 if r % 2 else 14000}"/>'))
        for i, side in enumerate(("lnL", "lnR", "lnT", "lnB")):  # 格線只留水平線；ln* 要排在填色前面
            fill = "<a:noFill/>" if i < 2 else hline
            tcPr.insert(i, parse_xml(f'<a:{side} {nsdecls("a")} w="{Pt(0.75)}">{fill}</a:{side}>'))
animate(s, fx(gf, "blinds", "after"))
notes(s, d["notes"])

# ─── 08 平均不是加總（K3，原生圖表） ───────────────────────
d = C.AVERAGE
s = new_slide(8, d["title"])
data = CategoryChartData()
data.categories = d["categories"]
for name, vals in d["series"]:
    data.add_series(name, vals)
cf = s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(0.5), Inches(1.5), Inches(5.8), Inches(3.35), data)
ch = cf.chart
ch.has_title = False
ch.font.size = Pt(11)
ch.font.color.rgb = RGBColor.from_string(MUTE)
cjk_font(ch.font, FONT)
ch.has_legend = True
ch.legend.position, ch.legend.include_in_layout = XL_LEGEND_POSITION.TOP, False
plot = ch.plots[0]
plot.gap_width, plot.overlap = 80, -10
plot.has_data_labels = True
dl = plot.data_labels
dl.font.size = Pt(12)
dl.font.color.rgb = RGBColor.from_string(INK)
dl.number_format, dl.number_format_is_linked, dl.position = "0.0", False, XL_LABEL_POSITION.OUTSIDE_END
for ser, (color, a) in zip(plot.series, ((INK, 0.7), (ACC, 1.0))):
    ser.format.fill.solid()
    ser.format.fill.fore_color.rgb = RGBColor.from_string(color)
    if a < 1:
        ser.format.fill._xPr.find(qn("a:solidFill"))[0].append(parse_xml(f'<a:alpha {nsdecls("a")} val="{round(a * 100000)}"/>'))
cat, val = ch.category_axis, ch.value_axis
cat.format.line.color.rgb = RGBColor.from_string(INK)
cat.format.line._get_or_add_ln().find(qn("a:solidFill"))[0].append(parse_xml(f'<a:alpha {nsdecls("a")} val="25000"/>'))
cat.major_tick_mark = val.major_tick_mark = XL_TICK_MARK.NONE
val.format.line.fill.background()
val.has_major_gridlines = True
val.major_gridlines.format.line.color.rgb = RGBColor.from_string(INK)
val.major_gridlines.format.line.width = Pt(0.75)
val.major_gridlines.format.line._get_or_add_ln().find(qn("a:solidFill"))[0].append(parse_xml(f'<a:alpha {nsdecls("a")} val="15000"/>'))
# 圖表區和繪圖區透明：python-pptx 沒有這兩個的填色 API
cs = ch._chartSpace
nofill = f'<c:spPr {nsdecls("c", "a")}><a:noFill/><a:ln><a:noFill/></a:ln></c:spPr>'
cs.find(qn("c:chart")).addnext(parse_xml(nofill))
cs.find(qn("c:chart")).find(qn("c:plotArea")).append(parse_xml(nofill))
text(s, 4.8, 1.5, 1.5, 0.25, d["tag"], CAPTION, MUTE, align=RIGHT)
card(s, 6.6, 1.55, 2.9, 3.3)
seq = [fx(cf, "wipe", "click", side="b", build="category")]
for i, (t, y) in enumerate(zip(d["points"], (1.8, 2.8, 3.8))):
    seq.append(fx([text(s, 6.85, y + 0.03, 0.35, 0.3, f"{i + 1:02d}", CAPTION, ACC, True),
                   text(s, 7.2, y, 2.1, 0.9, t, BODY, INK)], "fade", "after"))
animate(s, *seq)
notes(s, d["notes"])

# ─── 09 兩輪收斂（K1） ─────────────────────────────────────
d = C.FUNNEL
s = new_slide(9, d["title"])
seq = []
for i, (num, desc) in enumerate(d["kpis"]):
    x, hot = 0.5 + i * 3.2, i == 2
    cd = card(s, x, 1.6, 2.6, 2.5, "hot" if hot else "glass")
    big = text(s, x, 1.85, 2.6, 1.0, num, DISPLAY, ACC if hot else INK, True, CENTER, MID, fx="glow" if hot else "shadow")
    seq.append(fx([cd, big, text(s, x + 0.2, 3.0, 2.2, 0.9, desc, BODY, MUTE, align=CENTER)], "zoom", "after"))
    if i < 2:
        arrow = box(s, 3.175 + i * 3.2, 2.65, 0.45, 0.35, MSO_SHAPE.RIGHT_ARROW, fill=INK, fa=0.5)
        seq.append(fx(arrow, "wipe", "after", side="l"))
text(s, 0.5, 4.4, 9.0, 0.4, d["bottom"], BODY, INK, True)
animate(s, *seq, fx(big, "grow", "click", scale=115, back=True))
notes(s, d["notes"])

# ─── 10 隱私（G2，2×3 icon 格） ────────────────────────────
d = C.PRIVACY
s = new_slide(10, d["title"])
pngs = icons([c[0] for c in d["cells"]], hot=5)
seq, rest = [], []
for i, ((_, head, desc), png) in enumerate(zip(d["cells"], pngs)):
    x, y, hot = 0.5 + (i % 3) * 3.1, (1.55, 3.1)[i // 3], i == 5
    parts = [card(s, x, y, 2.8, 1.35, "hot" if hot else "glass"),
             s.shapes.add_picture(str(png), Inches(x + 0.25), Inches(y + 0.25), Inches(0.45), Inches(0.45)),
             text(s, x + 0.85, y + 0.27, 1.8, 0.4, head, H2, INK, True),
             text(s, x + 0.25, y + 0.78, 2.35, 0.5, desc, BODY, MUTE)]
    if hot:
        seq.append(fx(parts, "zoom", "click"))
    else:
        seq.append(fx(parts, "fade", "with", delay=i * 150))
        rest += parts
text(s, 0.5, 4.6, 9.0, 0.3, d["caption"], BODY, MUTE)
animate(s, *seq, fx(rest, "transparency", "with", to=0.5))
notes(s, d["notes"])

# ─── 11 重用（C2 變體，用形狀排成表） ──────────────────────
d = C.REUSE
s = new_slide(11, d["title"])
for t, x, w, color in zip(d["header"], (0.7, 2.9, 6.2), (2.0, 3.1, 3.2), (MUTE, ACC, MUTE)):
    text(s, x, 1.5, w, 0.4, t, CAPTION, color, True, anchor=MID)
seq = []
for (scene, ans_, back), y in zip(d["rows"], (1.95, 2.7, 3.45, 4.2)):
    seq.append(fx([card(s, 0.5, y, 9.0, 0.65), text(s, 0.7, y, 2.0, 0.65, scene, H2, INK, True, anchor=MID),
                   text(s, 2.9, y, 3.1, 0.65, ans_, BODY, INK, anchor=MID),
                   text(s, 6.2, y, 3.2, 0.65, back, BODY, MUTE, anchor=MID)], "fly", "click", side="l"))
animate(s, *seq)
notes(s, d["notes"])

# ─── 12 結論與行動（E1） ───────────────────────────────────
d = C.ACTIONS
s = new_slide(12, d["title"], footer=d["footer"])
seq = []
for i, ((t, page, short), y) in enumerate(zip(d["items"], (1.55, 2.65, 3.75))):
    # python-pptx 做不出 slide zoom，回看入口改成玻璃卡加跳頁超連結
    entry = label(card(s, 7.5, y, 1.8, 1.0125), [(f"第 {page} 頁", CAPTION, ACC, False), (short, H2, INK, True)])
    links.append((entry, page))
    seq.append(fx([badge(s, 0.5, y + 0.27, 0.45, i + 1), text(s, 1.15, y, 5.3, 1.0, t, BODY, INK, anchor=MID), entry],
                  "float", "click"))
animate(s, *seq)
notes(s, d["notes"])

for sp, page in links:
    sp.click_action.target_slide = slides[page - 1]

out = HERE / "JEV儀表板_python-pptx.pptx"
prs.save(out)

# ─── 自我檢查：頁數、備忘、轉場、每頁效果數 ────────────────
EXPECT = [5, 13, 22, 18, 15, 15, 1, 9, 12, 44, 16, 9]
assert len(slides) == 12 and all(x.notes_slide.notes_text_frame.text for x in slides)
for i, x in enumerate(slides, 1):
    el = x._element
    kind = "morph" if el.find("{http://schemas.openxmlformats.org/markup-compatibility/2006}AlternateContent") is not None else etree.QName(el.find(qn("p:transition"))[0]).localname
    assert kind == TRANSITION.get(i, "morph"), (i, kind)
    n = sum(1 for c in el.iter(qn("p:cTn")) if c.get("presetClass"))
    assert n == EXPECT[i - 1], (i, n)
    ids = [c.get("id") for c in el.iter(qn("p:cTn"))]
    assert len(ids) == len(set(ids)), (i, "cTn id 重複")
print(f"done: {out}  ({time.time() - T0:.1f}s)")
